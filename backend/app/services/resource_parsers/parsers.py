from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from pypdf import PdfReader


ALLOWED_EXTENSIONS = {".txt", ".csv", ".docx", ".pdf", ".xlsx", ".json", ".srt", ".vtt"}
MAX_SEGMENT_CHARS = 1500


class ResourceParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedSegment:
    locator_type: str
    locator_value: str
    text: str


def _clean_text(value: object) -> str:
    return re.sub(r"[ \t\f\v]+", " ", str(value or "")).strip()


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ResourceParseError("文本编码无法识别，仅支持 UTF-8、UTF-8-SIG 或 GB18030")


def _chunk(segment: ParsedSegment) -> list[ParsedSegment]:
    text = segment.text.strip()
    if len(text) <= MAX_SEGMENT_CHARS:
        return [segment]
    parts: list[ParsedSegment] = []
    start = 0
    part = 1
    while start < len(text):
        end = min(start + MAX_SEGMENT_CHARS, len(text))
        if end < len(text):
            split_at = max(text.rfind("。", start, end), text.rfind("\n", start, end))
            if split_at > start + MAX_SEGMENT_CHARS // 2:
                end = split_at + 1
        value = f"{segment.locator_value}#{part}" if segment.locator_value else str(part)
        parts.append(ParsedSegment(segment.locator_type, value, text[start:end].strip()))
        start = end
        part += 1
    return parts


def _finalize(segments: list[ParsedSegment]) -> list[ParsedSegment]:
    output: list[ParsedSegment] = []
    for segment in segments:
        cleaned = _clean_text(segment.text)
        if not cleaned:
            continue
        output.extend(_chunk(ParsedSegment(segment.locator_type, segment.locator_value, cleaned)))
    if not output:
        raise ResourceParseError("文件中没有可提取的文本")
    return output


def _parse_txt(data: bytes) -> list[ParsedSegment]:
    text = _decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    if len(paragraphs) == 1:
        paragraphs = [line.strip() for line in text.splitlines() if line.strip()]
    return [ParsedSegment("paragraph", str(index), paragraph) for index, paragraph in enumerate(paragraphs, 1)]


def _parse_csv(data: bytes) -> list[ParsedSegment]:
    text = _decode_text(data)
    rows = list(csv.reader(io.StringIO(text, newline="")))
    if not rows:
        return []
    header = [cell.strip().lower() for cell in rows[0]]
    candidates = {"text", "正文", "文本", "内容", "body", "content", "answer", "回答"}
    indexes = [index for index, name in enumerate(header) if name in candidates]
    data_rows = rows[1:] if indexes else rows
    segments: list[ParsedSegment] = []
    for row_number, row in enumerate(data_rows, 2 if indexes else 1):
        values = [row[index] for index in indexes if index < len(row)] if indexes else row
        text_value = " | ".join(_clean_text(value) for value in values if _clean_text(value))
        if text_value:
            segments.append(ParsedSegment("row", str(row_number), text_value))
    return segments


def _parse_docx(data: bytes) -> list[ParsedSegment]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ResourceParseError("DOCX 文件结构无效") from exc
    root = ElementTree.fromstring(xml)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    segments: list[ParsedSegment] = []
    for paragraph_number, paragraph in enumerate(root.iter(f"{namespace}p"), 1):
        text = "".join(node.text or "" for node in paragraph.iter(f"{namespace}t"))
        if text.strip():
            segments.append(ParsedSegment("paragraph", str(paragraph_number), text))
    return segments


def _parse_pdf(data: bytes) -> list[ParsedSegment]:
    try:
        reader = PdfReader(io.BytesIO(data))
        segments = [
            ParsedSegment("page", str(page_number), page.extract_text() or "")
            for page_number, page in enumerate(reader.pages, 1)
        ]
    except Exception as exc:
        raise ResourceParseError(f"PDF 解析失败：{exc}") from exc
    if not any(segment.text.strip() for segment in segments):
        raise ResourceParseError("PDF 没有可提取文本，可能需要先执行 OCR")
    return segments


def _parse_xlsx(data: bytes) -> list[ParsedSegment]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ResourceParseError(f"XLSX 解析失败：{exc}") from exc
    segments: list[ParsedSegment] = []
    try:
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                values = [_clean_text(value) for value in row if value is not None]
                if values:
                    segments.append(
                        ParsedSegment("sheet_row", f"{worksheet.title}!{row_number}", " | ".join(values))
                    )
    finally:
        workbook.close()
    return segments


def _parse_json(data: bytes) -> list[ParsedSegment]:
    try:
        payload = json.loads(_decode_text(data))
    except json.JSONDecodeError as exc:
        raise ResourceParseError(f"JSON 格式无效：第 {exc.lineno} 行") from exc
    segments: list[ParsedSegment] = []
    preferred = {"text", "正文", "文本", "内容", "body", "content", "answer", "回答"}

    def visit(value: object, path: str, preferred_context: bool = False) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                child_path = f"{path}.{key}" if path else f"$.{key}"
                visit(child, child_path, str(key).lower() in preferred)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                visit(child, f"{path}[{index}]", preferred_context)
        elif isinstance(value, str) and (preferred_context or len(value.strip()) >= 2):
            segments.append(ParsedSegment("json_path", path or "$", value))

    visit(payload, "$")
    return segments


_TIMECODE = re.compile(
    r"(?P<start>\d{1,2}:\d{2}(?::\d{2})?[,.]\d{3})\s*-->\s*(?P<end>\d{1,2}:\d{2}(?::\d{2})?[,.]\d{3})"
)


def _parse_subtitles(data: bytes) -> list[ParsedSegment]:
    text = _decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    segments: list[ParsedSegment] = []
    blocks = re.split(r"\n\s*\n", text)
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        time_index = next((index for index, line in enumerate(lines) if _TIMECODE.search(line)), None)
        if time_index is None:
            continue
        match = _TIMECODE.search(lines[time_index])
        caption = " ".join(lines[time_index + 1 :])
        if match and caption:
            locator = f"{match.group('start')}-->{match.group('end')}"
            segments.append(ParsedSegment("timestamp", locator, re.sub(r"<[^>]+>", "", caption)))
    return segments


def parse_resource(filename: str, data: bytes) -> list[ParsedSegment]:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        supported = "、".join(sorted(ALLOWED_EXTENSIONS))
        raise ResourceParseError(f"不支持 {extension or '无扩展名'} 文件，支持格式：{supported}")
    parser = {
        ".txt": _parse_txt,
        ".csv": _parse_csv,
        ".docx": _parse_docx,
        ".pdf": _parse_pdf,
        ".xlsx": _parse_xlsx,
        ".json": _parse_json,
        ".srt": _parse_subtitles,
        ".vtt": _parse_subtitles,
    }[extension]
    try:
        return _finalize(parser(data))
    except ResourceParseError:
        raise
    except Exception as exc:
        raise ResourceParseError(f"{extension} 文件解析失败：{exc}") from exc
